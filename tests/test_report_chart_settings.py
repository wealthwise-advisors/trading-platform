"""
The exported HTML report draws with the live chart's settings.

Verification found the report drawing fixed defaults of its own: all four
oscillator rows whatever the chart had switched off, and the Volume Profile
levels in factory colours with no bubbles, titles or names. The chart now sends
its settings with Export Report (api/schemas/chart_settings.py) and the report
draws with them (api/report/chart_settings_draw.py). These tests pin each
setting's effect on the report's figure, the contract both sides share, and the
plumbing through the endpoint for all five formats.
"""

import io
import itertools
import json
import math
import re
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from pydantic import ValidationError

from api.report.chart_settings_draw import (
    MAX_LEVEL_MARKERS, compute_volume_profile, compute_volume_profiles,
    js_number, js_to_fixed, oscillator_row_heights,
)
from api.report.report import _candlestick_chart
from api.schemas.chart_settings import ChartSettings
from src.analysis.indicators import calc_volume_profile

FIXTURE = Path(__file__).parent / "fixtures" / "chart_settings_factory.json"
OSC_KEYS = ["rsi2", "stochrsi", "rsi13", "mfi"]
OSC_TITLES = {"rsi2": "RSI(2)", "stochrsi": "StochRSI", "rsi13": "RSI(13)", "mfi": "MFI"}
OSC_TRACES = {"rsi2": ["RSI(2)"], "stochrsi": ["FullK", "FullD"], "rsi13": ["RSI(13)"], "mfi": ["MoneyFlowIndex"]}
OSC_LEVELS = {"rsi2": [2, 94], "stochrsi": [20, 80], "rsi13": [30, 70], "mfi": [20, 80]}
VP_NAMES = {"Volume Profile", "POC", "VAHigh", "VALow", "ProfileHigh", "ProfileLow"}


# ── fixtures and helpers ─────────────────────────────────────────────────────

def _df(bars=400, start="2025-01-02 09:30", freq="5min", seed=3, volume=True):
    idx = pd.date_range(start, periods=bars, freq=freq)
    rng = np.random.default_rng(seed)
    close = np.round((5000 + rng.normal(0, 2, bars).cumsum()) * 4) / 4
    df = pd.DataFrame({
        "open": close, "close": close,
        "high": close + rng.integers(1, 8, bars) * 0.25,
        "low": close - rng.integers(1, 8, bars) * 0.25,
    }, index=idx)
    if volume:
        df["volume"] = rng.integers(100, 900, bars).astype(float)
    return df


def _sessions(days=("2025-01-02", "2025-01-03", "2025-01-06")):
    return pd.concat([_df(bars=78, start=f"{d} 09:30", seed=10 + i) for i, d in enumerate(days)])


def _fig(chart=None, df=None):
    df = _df() if df is None else df
    results = SimpleNamespace(price_data=df, trades=[], symbol="ES", strategy_name="probe")
    return _candlestick_chart(results, chart_settings=None if chart is None else ChartSettings.model_validate(chart))


def _vp(**fields):
    return {"volumeProfile": fields}


def _plot(**fields):
    return {"show": True, "values": "numerical", "drawAs": "line", "style": "solid", "width": 1,
            "color": "#38bdf8", "bubble": True, "title": True, **fields}


def _traces(fig, name):
    return [t for t in fig.data if t.name == name]


def _row_titles(fig):
    layout = fig.layout.to_plotly_json()
    keys = sorted((k for k in layout if re.fullmatch(r"yaxis\d*", k)), key=lambda k: int(k[5:] or 1))
    return [layout[k]["title"]["text"] for k in keys]


def _level_lines(fig):
    out = {}
    for s in fig.layout.shapes:
        if s.type == "line" and str(s.xref).endswith("domain") and s.y0 == s.y1:
            out.setdefault(s.yref, []).append(s.y0)
    return {k: sorted(v) for k, v in out.items()}


def _vp_annotations(fig):
    # Swing headers are pinned to the paper; everything the study adds is on y.
    return [a for a in fig.layout.annotations if a.yref in ("y", "y domain")]


def _bubbles(fig):
    return [a for a in _vp_annotations(fig) if a.bgcolor and a.bgcolor.startswith("#")]


@pytest.fixture(scope="module")
def base_df():
    return _df()


# ── the contract both sides share ────────────────────────────────────────────

def test_factory_defaults_are_the_shared_fixture():
    fixture = json.loads(FIXTURE.read_text(encoding="utf-8"))
    assert ChartSettings().model_dump(by_alias=True) == fixture
    assert ChartSettings.model_validate(fixture) == ChartSettings()


@pytest.mark.parametrize("bad", [
    _vp(plots={"poc": _plot(color="red")}),
    _vp(plots={"poc": _plot(style="zigzag")}),
    _vp(plots={"poc": _plot(width=9)}),
    _vp(plots={"poc": _plot(drawAs="histogram")}),
    _vp(bins=0),
    _vp(timePer="MONTH"),
    {"oscillators": {"stoch": True}},
    {"surprise": 1},
], ids=["colour", "style", "width", "draw-as", "bins", "time-per", "unknown-oscillator", "unknown-field"])
def test_settings_the_chart_could_not_have_sent_are_refused(bad):
    with pytest.raises(ValidationError):
        ChartSettings.model_validate(bad)


def test_the_value_area_edges_are_reconciled_as_on_the_web():
    s = ChartSettings.model_validate(_vp(plots={"vah": _plot(show=True, style="dash"),
                                                "val": _plot(show=False, style="dash")}))
    assert (s.volume_profile.plots.vah.show, s.volume_profile.plots.val.show) == (True, True)


# ── oscillators: exactly the rows switched on ────────────────────────────────

def test_no_settings_draws_what_the_chart_shows_untouched(base_df):
    fig = _fig(df=base_df)
    assert _row_titles(fig) == ["Price", "RSI(2)", "StochRSI", "RSI(13)"]
    assert not _traces(fig, "MoneyFlowIndex")


@pytest.mark.parametrize("combo", list(itertools.product([False, True], repeat=4)),
                         ids=lambda c: "+".join(k for k, on in zip(OSC_KEYS, c) if on) or "none")
def test_the_report_draws_exactly_the_oscillators_switched_on(combo, base_df):
    on = [k for k, v in zip(OSC_KEYS, combo) if v]
    fig = _fig({"oscillators": dict(zip(OSC_KEYS, combo))}, base_df)

    assert _row_titles(fig) == ["Price"] + [OSC_TITLES[k] for k in on]
    for key in OSC_KEYS:
        for name in OSC_TRACES[key]:
            axes = [t.yaxis for t in _traces(fig, name)]
            if key in on:
                assert axes == [f"y{on.index(key) + 2}"], f"{name} on {axes}"
            else:
                assert axes == [], f"{name} drawn although {key} is off"
    assert _level_lines(fig) == {f"y{on.index(k) + 2}": OSC_LEVELS[k] for k in on}

    swing_rows = {t.yaxis for t in fig.data if getattr(t, "mode", None) == "markers+text" and t.yaxis != "y"}
    assert swing_rows == {f"y{i + 2}" for i in range(len(on))}, swing_rows


def test_row_heights_match_the_chart():
    assert oscillator_row_heights(0) == [1.0]
    assert oscillator_row_heights(2) == pytest.approx([0.68, 0.16, 0.16])
    fig = _fig({"oscillators": dict.fromkeys(OSC_KEYS, False)})
    assert list(fig.layout.yaxis.domain) == [0, 1]


# ── Volume Profile: each setting reaches the report ──────────────────────────

def test_colour_width_and_style(base_df):
    [poc] = _traces(_fig(_vp(plots={"poc": _plot(color="#ff8800", width=3, style="longdash")}), base_df), "POC")
    assert (poc.mode, poc.line.color, poc.line.width, poc.line.dash) == ("lines", "#ff8800", 3, "longdash")


@pytest.mark.parametrize("draw_as,symbol", [("points", "circle"), ("squares", "square"), ("triangles", "triangle-up")])
def test_draw_as_markers(draw_as, symbol, base_df):
    fig = _fig(_vp(plots={"vah": _plot(drawAs=draw_as, width=2, style="dash", color="#7dd3fc")}), base_df)
    [vah] = _traces(fig, "VAHigh")
    assert (vah.mode, vah.marker.symbol, vah.marker.size) == ("markers", symbol, 7)
    assert len(vah.x) <= MAX_LEVEL_MARKERS + 1


def test_show_value_area_off_removes_both_edges_everywhere(base_df):
    off = dict(show=False, style="dash", color="#7dd3fc")
    fig = _fig(_vp(plots={"vah": _plot(**off), "val": _plot(**off)}), base_df)
    names = {t.name for t in fig.data}
    assert "POC" in names and not {"VAHigh", "VALow"} & names
    assert all("VAHigh" not in a.text and "VALow" not in a.text for a in _vp_annotations(fig))
    assert fig.layout.title.text.endswith("Yes, No, 70, 50)</span>")


def test_plot_names_follow_the_switch(base_df):
    named = [a for a in _vp_annotations(_fig(df=base_df)) if a.text in ("POC", "VAHigh", "VALow")]
    assert {a.text: a.font.color for a in named} == {"POC": "#38bdf8", "VAHigh": "#7dd3fc", "VALow": "#7dd3fc"}
    assert len(named) == 3
    off = _fig(_vp(showPlotNames=False), base_df)
    assert not [a for a in _vp_annotations(off) if a.text in ("POC", "VAHigh", "VALow")]


def test_bubbles_per_plot_and_on_the_price_axis_side(base_df):
    fig = _fig(df=base_df)
    assert len(_bubbles(fig)) == 3
    assert all((b.x, b.xanchor) == (0, "right") for b in _bubbles(fig))
    assert fig.layout.yaxis.side == "left"

    assert len(_bubbles(_fig(_vp(plots={"poc": _plot(bubble=False)}), base_df))) == 2

    right = _fig(_vp(leftAxis=False), base_df)
    assert all((b.x, b.xanchor) == (1, "left") for b in _bubbles(right))
    assert right.layout.yaxis.side == "right"


def test_a_bubble_prints_its_level_to_the_cent_as_the_chart_does(base_df):
    fig = _fig(df=base_df)
    [poc] = _traces(fig, "POC")
    [bubble] = [b for b in _bubbles(fig) if b.bgcolor == "#38bdf8"]
    assert bubble.y == poc.y[0]
    assert bubble.text == js_to_fixed(poc.y[0], 2)


def test_the_title_line_lists_titled_plots_in_tab_order(base_df):
    [line] = [a.text for a in _vp_annotations(_fig(df=base_df)) if a.text.startswith("POC: ")]
    assert re.fullmatch(r"POC: \d+\.\d\d  ·  VAHigh: \d+\.\d\d  ·  VALow: \d+\.\d\d", line), line

    untitled = _fig(_vp(plots={"poc": _plot(title=False)}), base_df)
    [line] = [a.text for a in _vp_annotations(untitled) if "VAHigh: " in a.text]
    assert re.fullmatch(r"VAHigh: \d+\.\d\d  ·  VALow: \d+\.\d\d", line), line


def test_input_names_follow_the_switch_and_print_as_the_chart_prints_them(base_df):
    fig = _fig(_vp(rowMode="MANUAL", rowHeight=0.25, multiplier=2, valueArea=65, opacity=35), base_df)
    assert fig.layout.title.text.endswith(
        "VolumeProfile(MANUAL, 0.25, CHART, 2, Yes, 1000, Yes, Yes, 65, 35)</span>")
    assert "VolumeProfile(" not in _fig(_vp(showInputNames=False), base_df).layout.title.text


def test_opacity_sets_the_histogram_colours(base_df):
    [bar] = _traces(_fig(_vp(opacity=80), base_df), "Volume Profile")
    assert set(bar.marker.color) == {"rgba(56,189,248,0.544)", "rgba(56,189,248,0.208)"}


def test_the_histogram_sits_on_the_overlay_axis_not_the_date_axis(base_df):
    fig = _fig(df=base_df)
    [bar] = _traces(fig, "Volume Profile")
    assert (bar.xaxis, bar.yaxis) == ("x9", "y")
    assert fig.layout.xaxis9.overlaying == "x"


def test_bins_and_value_area_change_the_levels(base_df):
    fig = _fig(_vp(bins=96, valueArea=50), base_df)
    prof = compute_volume_profile(base_df["high"].tolist(), base_df["low"].tolist(),
                                  base_df["volume"].tolist(), 96, 0.5)
    assert len(_traces(fig, "Volume Profile")[0].y) == 96
    assert _traces(fig, "POC")[0].y[0] == prof.poc
    assert _traces(fig, "VAHigh")[0].y[0] == prof.vah
    assert _traces(fig, "VALow")[0].y[0] == prof.val


def test_profile_high_and_low_when_shown(base_df):
    edge = dict(show=True, style="dot", color="#94a3b8")
    fig = _fig(_vp(plots={"profileHigh": _plot(**edge), "profileLow": _plot(**edge)}), base_df)
    [bar] = _traces(fig, "Volume Profile")
    half = bar.width / 2
    assert _traces(fig, "ProfileHigh")[0].y[0] == bar.y[-1] + half
    assert _traces(fig, "ProfileLow")[0].y[0] == bar.y[0] - half


def test_manual_row_height_sets_the_bin_count(base_df):
    [bar] = _traces(_fig(_vp(rowMode="MANUAL", rowHeight=0.5), base_df), "Volume Profile")
    expected = int(np.floor((base_df["high"].max() - base_df["low"].min()) / 0.5 + 0.5))
    assert len(bar.y) == expected


def test_day_profiles_are_drawn_inside_their_own_sessions():
    fig = _fig(_vp(timePer="DAY", plots={"poc": _plot(color="#ff8800", width=3)}), _sessions())
    assert not VP_NAMES & {t.name for t in fig.data}, "the whole-chart overlay should not draw"
    rects = [s for s in fig.layout.shapes if s.type == "rect" and s.yref == "y"]
    lines = [s for s in fig.layout.shapes if s.type == "line" and s.xref == "x"]
    assert {pd.Timestamp(s.x0).date() for s in rects} == {pd.Timestamp(d).date() for d in ("2025-01-02", "2025-01-03", "2025-01-06")}
    assert len(lines) == 9, "POC, VAHigh and VALow for each of three sessions"
    poc_lines = [s for s in lines if s.line.color == "#ff8800"]
    assert len(poc_lines) == 3 and all(s.line.width == 3 for s in poc_lines)
    # The chart writes naive boundaries with toNaiveString, which has no
    # milliseconds; the report must land on the same whole second.
    assert all(pd.Timestamp(s.x1) == pd.Timestamp(s.x1).floor("s") for s in rects + lines)


def test_multiplier_and_profile_count():
    df = _sessions()
    two = ChartSettings.model_validate(_vp(timePer="DAY", multiplier=2)).volume_profile
    assert [len({s.start.date(), s.end.date()}) for s in compute_volume_profiles(df, two)] == [2, 1]
    last = ChartSettings.model_validate(_vp(timePer="DAY", maxProfiles=1)).volume_profile
    [only] = compute_volume_profiles(df, last)
    assert only.start.date() == pd.Timestamp("2025-01-06").date()


def test_days_are_split_by_the_date_written_on_each_bar():
    idx = pd.to_datetime(["2025-01-02 19:30", "2025-01-02 20:30", "2025-01-02 23:55", "2025-01-03 00:05"])
    df = pd.DataFrame({"open": [1.0, 2, 3, 4], "high": [2.0, 3, 4, 5], "low": [0.5, 1, 2, 3],
                       "close": [1.0, 2, 3, 4], "volume": [10.0, 10, 10, 10]}, index=idx)
    vp = ChartSettings.model_validate(_vp(timePer="DAY")).volume_profile
    assert [(s.start.hour, s.end.hour) for s in compute_volume_profiles(df, vp)] == [(19, 23), (0, 0)]


@pytest.mark.parametrize("hidden", [{"on": False}, {"showStudy": False}, {"onExpansion": False}],
                         ids=["off", "show-study-off", "on-expansion-off"])
def test_nothing_is_drawn_when_the_study_is_hidden(hidden, base_df):
    fig = _fig(_vp(**hidden), base_df)
    assert not VP_NAMES & {t.name for t in fig.data}
    assert not _vp_annotations(fig)
    # Input names follow the chart: shown unless the study itself is off or hidden.
    assert ("VolumeProfile(" in fig.layout.title.text) == (hidden == {"onExpansion": False})


def test_without_volume_no_profile_is_drawn():
    fig = _fig(df=_df(volume=False))
    assert not VP_NAMES & {t.name for t in fig.data}


# ── the port itself ──────────────────────────────────────────────────────────

GOLDEN = Path(__file__).parent / "fixtures" / "volume_profile_golden.json"


@pytest.fixture(scope="module")
def golden():
    return json.loads(GOLDEN.read_text(encoding="utf-8"))


def _as_browser(p):
    return {"prices": p.prices, "volumes": p.volumes, "poc": p.poc, "val": p.val, "vah": p.vah, "binSize": p.bin_size}


def test_the_port_reproduces_the_browsers_profile_exactly(golden):
    # Expected values were computed by web/src/lib/volumeProfile.ts, which
    # web/src/lib/volumeProfile.test.ts pins to the same file. Exact equality,
    # not a tolerance: a bubble prints these to the cent.
    bars = golden["bars"]
    highs, lows = [b["h"] for b in bars], [b["l"] for b in bars]
    vols = [0.0 if b["v"] is None else float(b["v"]) for b in bars]
    for case in golden["single"]:
        got = compute_volume_profile(highs, lows, vols, case["bins"], case["valueArea"] / 100)
        assert _as_browser(got) == case["expected"], case["bins"]


def test_the_port_reproduces_the_browsers_session_profiles_exactly(golden):
    bars = golden["bars"]
    df = pd.DataFrame({"high": [b["h"] for b in bars], "low": [b["l"] for b in bars],
                       "volume": [np.nan if b["v"] is None else float(b["v"]) for b in bars]},
                      index=pd.to_datetime([b["t"] for b in bars]))
    for case in golden["multi"]:
        o = case["options"]
        vp = ChartSettings.model_validate(_vp(
            timePer=o.get("timePer", "CHART"), multiplier=o.get("multiplier", 1),
            maxProfiles=o.get("maxProfiles", 1000), rowMode=o.get("rowMode", "AUTOMATIC"),
            rowHeight=o.get("customRowHeight", 1), bins=o.get("bins", 48), valueArea=o.get("valueArea", 70),
        )).volume_profile
        got = [{"startT": s.start.isoformat(), "endT": s.end.isoformat(), "profile": _as_browser(s.profile)}
               for s in compute_volume_profiles(df, vp)]
        assert got == case["expected"], o


def test_levels_print_to_the_cent_exactly_as_the_browser_prints_them(golden):
    assert len(golden["fixed"]) > 100
    mismatches = [(v, d, js, js_to_fixed(v, d)) for v, d, js in golden["fixed"] if js_to_fixed(v, d) != js]
    assert not mismatches, mismatches[:5]


def test_the_old_server_profile_is_not_what_the_chart_draws(golden):
    # Why the report no longer uses calc_volume_profile: its floor division
    # (a // w) files some bar edges one bucket lower than the browser's
    # Math.floor(a / w), so its POC and value area could differ from the chart.
    # Recorded so the reason survives; see chart_settings_draw.py's docstring.
    bars = golden["bars"]
    lo = min(b["l"] for b in bars)
    hi = max(b["h"] for b in bars)
    w = (hi - lo) / 48
    edges = [b[k] for b in bars for k in ("h", "l")]
    assert any(math.floor((x - lo) / w) != int((x - lo) // w) for x in edges)
    assert calc_volume_profile is not None


@pytest.mark.parametrize("value,digits,text", [
    (100.125, 2, "100.13"), (5000.375, 2, "5000.38"), (100.0, 2, "100.00"), (1.005, 2, "1.00"),
    (0.8 * 0.68, 3, "0.544"), (-0.001, 2, "-0.00"), (-0.0, 2, "0.00"),
])
def test_js_to_fixed_matches_number_to_fixed(value, digits, text):
    assert js_to_fixed(value, digits) == text


@pytest.mark.parametrize("value,text", [
    (1.0, "1"), (0.25, "0.25"), (70, "70"), (1e-05, "0.00001"), (0.1 + 0.2, "0.30000000000000004"),
])
def test_js_number_matches_a_template_literal(value, text):
    assert js_number(value) == text


# ── through the endpoint, every format ───────────────────────────────────────

@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient
    from api.main import app
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def backtest_id(client):
    r = client.post("/api/backtests", json={
        "symbol": "ES", "strategy_id": "ma_crossover", "timeframe": "5m",
        "start_date": "2024-01-01", "end_date": "2024-01-10", "data_source": "synthetic",
        "initial_capital": 100000, "params": {"fast": 9, "slow": 21},
    })
    assert r.status_code == 200, r.text
    return r.json()["backtest_id"]


def _report(client, backtest_id, fmt, chart=None):
    params = {"format": fmt}
    if chart is not None:
        params["chart"] = chart if isinstance(chart, str) else json.dumps(chart)
    return client.get(f"/api/backtests/{backtest_id}/report", params=params)


def _has_trace(html, name):
    return re.search(r'"name":\s*"' + re.escape(name) + '"', html) is not None


CHANGED = {
    "oscillators": {"rsi2": False, "stochrsi": True, "rsi13": False, "mfi": True},
    "volumeProfile": {"plots": {"poc": _plot(color="#ff8800")}, "showPlotNames": False},
}


def test_the_html_report_is_drawn_with_the_settings_sent(client, backtest_id):
    r = _report(client, backtest_id, "html", CHANGED)
    assert r.status_code == 200, r.text[:300]
    html = r.text
    assert _has_trace(html, "FullK") and _has_trace(html, "MoneyFlowIndex")
    assert not _has_trace(html, "RSI(2)") and not _has_trace(html, "RSI(13)")
    assert "#ff8800" in html


def test_the_html_report_without_settings_is_the_chart_untouched(client, backtest_id):
    html = _report(client, backtest_id, "html").text
    assert _has_trace(html, "RSI(2)") and _has_trace(html, "RSI(13)")
    assert not _has_trace(html, "MoneyFlowIndex")


@pytest.mark.parametrize("fmt", ["html", "csv", "xlsx", "pdf", "docx"])
@pytest.mark.parametrize("chart", [_vp(plots={"poc": _plot(color="red")}), "{not json"], ids=["bad-colour", "not-json"])
def test_invalid_settings_are_refused_for_every_format(client, backtest_id, fmt, chart):
    r = _report(client, backtest_id, fmt, chart)
    assert r.status_code == 400, r.text[:200]
    assert r.json()["detail"].startswith("chart settings are not valid")


def _document_text(fmt, content):
    if fmt == "csv":
        return content.decode("utf-8")
    if fmt == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content))
        return "\n".join(
            ws.title + ":" + "|".join("" if v is None else str(v) for row in ws.iter_rows(values_only=True) for v in row)
            for ws in wb.worksheets)
    if fmt == "docx":
        from docx import Document
        doc = Document(io.BytesIO(content))
        return "\n".join([p.text for p in doc.paragraphs]
                         + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    # ReportLab writes each page's content as an ASCII85 + Flate stream, and the
    # text in it as (string) Tj operators. Decoded here rather than with a PDF
    # library, which the project does not otherwise need.
    import base64
    import zlib
    text = []
    for stream in re.findall(rb"/Filter \[ /ASCII85Decode /FlateDecode \][^>]*>>\s*stream\r?\n(.*?)endstream",
                             content, re.S):
        # ReportLab ends the ASCII85 data with its "~>" terminator.
        raw = zlib.decompress(base64.a85decode(stream.strip(), adobe=True, ignorechars=b" \t\n\r\x0b"))
        text += [s.decode("latin-1") for s in re.findall(rb"\(((?:\\.|[^\\)])*)\)\s*Tj", raw)]
    return "\n".join(text)


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "pdf", "docx"])
def test_table_formats_carry_no_chart_and_the_settings_change_nothing_in_them(client, backtest_id, fmt):
    plain = _report(client, backtest_id, fmt)
    changed = _report(client, backtest_id, fmt, CHANGED)
    assert plain.status_code == changed.status_code == 200
    text = _document_text(fmt, plain.content)
    assert "Total Trades" in text or "total_trades" in text.lower(), "could not read the document's content"
    assert text == _document_text(fmt, changed.content)
    # Oscillator and profile names, not bare "RSI": a strategy such as
    # "RSI Divergence" is printed by name in every format.
    for term in ("RSI(2)", "RSI(13)", "StochRSI", "FullK", "MoneyFlowIndex", "MFI",
                 "Volume Profile", "POC", "VAHigh", "VALow"):
        assert term not in text, f"{fmt} mentions {term}"
